import os
import mimetypes
import magic
from io import BytesIO
import logging
import tiktoken
#imports for processing different filetypes 
import PyPDF2 #PDF
import docx #Word Docs
from pptx import Presentation #PowerPoint
from striprtf.striprtf import rtf_to_text

# Mapping of file extensions to MIME types
extension_to_mime = {
    '.pdf': 'application/pdf',
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.txt': 'text/plain',
    '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    '.rtf': 'application/rtf',
    # Add more mappings as needed
}

# Mapping of MIME types to extraction functions
mime_type_to_extractor = {
    'application/pdf': extract_text_from_pdf,
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document': extract_text_from_docx,
    'text/plain': extract_text_from_txt,
    'application/vnd.openxmlformats-officedocument.presentationml.presentation': extract_text_from_pptx,
    'application/rtf': extract_text_from_rtf,
    # Add more mappings as needed
}

def get_file_extension(filename):
    """
    Extracts file extension from filename
    """
    _, ext = os.path.splitext(filename)
    return ext.lower()

def detect_mime_type(file_bytes):
    """
    Detects MIME type of a file using magic numbers
    """
    try:
        mime = magic.from_buffer(file_bytes, mime=True)
        return mime
    except Exception as e:
        #logging.error(f"Error detecting MIME type: {e}")
        return None

def clean_text(text):
    """
    Cleans text by removing unwanted characters or formatting
    """
    #Removes zero-width spaces
    text = text.replace('\u200b', '')
    return text

def chunk_text(text, max_tokens=1000):
    """
    Splits text into chunks suitable for embeddings based on token count
    """
    text = clean_text(text)
    tokenizer = tiktoken.get_encoding("cl100k_base")#may need to change based on model
    tokens = tokenizer.encode(text)
    chunks = []
    for i in range(0, len(tokens), max_tokens):
        chunk_tokens = tokens[i:i + max_tokens]
        chunk = tokenizer.decode(chunk_tokens)
        chunks.append(chunk)
    return chunks


def extract_text_from_pdf(file_bytes):
    """
    Extracts text from a PDF file
    """
    try:
        pdf_reader = PyPDF2.PdfReader(BytesIO(file_bytes))
        text = ''
        for page in pdf_reader.pages:
            text += page.extract_text() or ''
        return text
    except Exception as e:
        logging.error(f"Error extracting text from PDF: {e}")
        return ''
    
def extract_text_from_docx(file_bytes):
    """
    Extracts text from a DOCX file
    """
    try:
        document = docx.Document(BytesIO(file_bytes))
        text = '\n'.join([para.text for para in document.paragraphs])
        return text
    except Exception as e:
        logging.error(f"Error extracting text from DOCX: {e}")
        return ''
    
def extract_text_from_txt(file_bytes):
    """
    Extracts text from a TXT file
    """
    try:
        text = file_bytes.decode('utf-8', errors='ignore')
        return text
    except Exception as e:
        logging.error(f"Error extracting text from TXT: {e}")
        return ''
    
def extract_text_from_pptx(file_bytes):
    """
    Extracts text from a PPTX file.
    """
    try:
        presentation = Presentation(BytesIO(file_bytes))
        text_runs = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return '\n'.join(text_runs)
    except Exception as e:
        logging.error(f"Error extracting text from PPTX: {e}")
        return ''

def extract_text_from_rtf(file_bytes):
    """
    Extracts text from an RTF file.
    """
    try:
        text = rtf_to_text(file_bytes.decode('utf-8', errors='ignore'))
        return text
    except Exception as e:
        logging.error(f"Error extracting text from RTF: {e}")
        return ''
    
def process_file(file_bytes, filename):
    """
    Main function to process a file: detects file type, extracts text, and chunks it.
    """
    extension = get_file_extension(filename)
    expected_mime = extension_to_mime.get(extension)
    detected_mime = detect_mime_type(file_bytes)

    if expected_mime != detected_mime:
        if expected_mime is not None:
            logging.warning(f"MIME type mismatch for file '{filename}': Expected '{expected_mime}', Detected '{detected_mime}'")
        mime_type = detected_mime or expected_mime
    else:
        mime_type = expected_mime

    extractor = mime_type_to_extractor.get(mime_type)
    
    if extractor:
        logging.info(f"Using extractor for MIME type: {mime_type}")
        text = extractor(file_bytes)
    else:
        error_message = f"No extractor found for MIME type '{mime_type}' or extension '{extension}'."
        logging.error(error_message)
        # Gracefully handle unsupported file types
        return None

    if not text:
        error_message = f"Could not extract text from file '{filename}'."
        logging.error(error_message)
        raise ValueError(error_message)

    chunks = chunk_text(text)
    return chunks
